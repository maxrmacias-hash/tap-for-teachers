from docx import Document
from docx.shared import Pt
d = Document(); d.add_heading("Lincoln Middle School — Lesson Plan", 0)
t = d.add_table(rows=4, cols=4); t.style = "Table Grid"
for i, (a, b) in enumerate([("Teacher", "Date"), ("Subject", "Grade"), ("Lesson Title", "Duration"), ("Unit", "Day / Sequence")]):
    t.cell(i, 0).text = a; t.cell(i, 2).text = b
    for c in (0, 2): t.cell(i, c).paragraphs[0].runs[0].bold = True
for h in ["Learning Objectives", "Standards Alignment", "Materials", "Procedure & Timing", "Differentiation", "Assessment / Exit Ticket", "Instructional Tips"]:
    d.add_heading(h, 2); d.add_paragraph("")
p = d.add_paragraph(); p.add_run("Reflection:")
d.save("samples/sample_template.docx")
