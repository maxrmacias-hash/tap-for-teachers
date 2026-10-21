"""PDF support: text extraction for past lessons, and best-effort PDF -> blank DOCX template rebuild."""
import io, re, subprocess, tempfile, os
import pdfplumber
from docx import Document
from docx.shared import Pt

def pdf_text(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f: f.write(data); fn = f.name
    try:
        return subprocess.run(["pdftotext", "-layout", fn, "-"], capture_output=True, text=True).stdout
    except FileNotFoundError:
        with pdfplumber.open(io.BytesIO(data)) as pdf: return "\n".join(p.extract_text() or "" for p in pdf.pages)
    finally: os.unlink(fn)

def _labels_only(text):
    """Keep only label lines (ending ':' or 'Label: value' -> 'Label:'); drop filled content."""
    out = []
    for line in (text or "").split("\n"):
        l = line.strip().replace("\u200b", "")
        m = re.match(r"^([A-Z][^:]{0,50}:)", l)
        if m: out.append(m.group(1))
    return out

def pdf_to_template(data: bytes) -> bytes:
    d = Document(); d.styles["Normal"].font.size = Pt(9)
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            tables = page.find_tables()
            bboxes = [t.bbox for t in tables]
            outside = page.filter(lambda o: not any(b[0] <= o.get("x0", 0) <= b[2] and b[1] <= o.get("top", 0) <= b[3] for b in bboxes))
            for line in (outside.extract_text() or "").split("\n"):
                l = line.strip()
                if not l: continue
                labs = _labels_only(l)
                if l.isupper() or len(l) < 40 and not labs: d.add_heading(l.title() if l.isupper() else l, 2)
                elif labs: d.add_paragraph("   ".join(labs))
            for t in tables:
                rows = t.extract()
                if not rows: continue
                n = max(len(r) for r in rows); tbl = d.add_table(rows=len(rows), cols=n); tbl.style = "Table Grid"
                for ri, r in enumerate(rows):
                    for ci, v in enumerate(r):
                        v = (v or "").replace("\u200b", "")
                        if ri == 0: keep = v.strip()
                        elif ci == 0 and len(v) < 60:                       # section names stay, drop times
                            keep = "\n".join(x for x in v.split("\n") if not re.search(r"\d+\s*min", x, re.I))
                        else: keep = "\n".join(_labels_only(v))
                        cell = tbl.cell(ri, ci); lines = [x for x in keep.split("\n") if x.strip()] or [""]
                        cell.paragraphs[0].text = lines[0]
                        for x in lines[1:]: cell.add_paragraph(x)
    out = io.BytesIO(); d.save(out); return out.getvalue()
