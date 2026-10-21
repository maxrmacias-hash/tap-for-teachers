"""Detect sections in a DOCX template and populate them preserving formatting."""
import io, re, copy
from docx import Document

LABEL_RE = re.compile(r"^\s*([A-Za-z][A-Za-z0-9 /&()'\-]{1,60}?)\s*:\s*(.*)$")

def _is_heading(p):
    s = (p.style.name or "").lower() if p.style is not None else ""
    if s.startswith("heading") or s == "title":
        return True
    t = p.text.strip()
    return bool(t) and len(t) < 60 and all(r.bold for r in p.runs if r.text.strip()) and p.runs != []

def detect_fields(doc):
    """Return list of fields: dict(id, label, kind, ref)."""
    fields, seen = [], set()
    def add(label, kind, ref):
        key = label.strip().rstrip(":").strip()
        if not key or key.lower() in seen: return
        seen.add(key.lower()); fields.append({"id": len(fields), "label": key, "kind": kind, "ref": ref})
    # tables: label cell + empty value cell (same row) or header row + empty row below
    for ti, tbl in enumerate(doc.tables):
        for ri, row in enumerate(tbl.rows):
            cells = row.cells
            for ci in range(len(cells) - 1):
                lab, val = cells[ci].text.strip(), cells[ci + 1].text.strip()
                if lab and not val and cells[ci]._tc is not cells[ci+1]._tc:
                    add(lab, "table_cell", (ti, ri, ci + 1))
            if ri + 1 < len(tbl.rows):
                below = tbl.rows[ri + 1].cells
                for ci, c in enumerate(cells):
                    if c.text.strip() and ci < len(below) and not below[ci].text.strip() and not any(x.text.strip() for x in below):
                        add(c.text, "table_below", (ti, ri + 1, ci))
    # label paragraphs inside cells ("Topic:", "Criteria for Success:") -> fill right after the label
    for ti, tbl in enumerate(doc.tables):
        seen_tc = set()
        for ri, row in enumerate(tbl.rows):
            for ci, cell in enumerate(row.cells):
                if id(cell._tc) in seen_tc: continue
                seen_tc.add(id(cell._tc))
                for pi, par in enumerate(cell.paragraphs):
                    t = par.text.strip()
                    if t.endswith(":") and len(t) < 60:
                        add(t, "cell_para", (ti, ri, ci, pi))
    # section rows of header tables: label in col 0, empty content cells -> "<row> | <column header>"
    for ti, tbl in enumerate(doc.tables):
        if len(tbl.rows) < 2: continue
        hdr = [c.text.strip() for c in tbl.rows[0].cells]
        if not all(hdr) or len(set(hdr)) < 2: continue
        for ri, row in enumerate(tbl.rows[1:], 1):
            cells = row.cells; lab = " / ".join(x for x in cells[0].text.split("\n") if x.strip())
            if not lab: continue
            uniq = []
            for ci, c in enumerate(cells[1:], 1):
                if c._tc is not cells[0]._tc and all(c._tc is not u[1]._tc for u in uniq): uniq.append((ci, c))
            for ci, c in uniq:
                if not c.text.strip():
                    add(lab if len(uniq) == 1 else f"{lab} | {hdr[ci]}", "table_cell", (ti, ri, ci))
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        t = p.text.strip()
        if not t: continue
        m = LABEL_RE.match(t)
        if m and not m.group(2).strip() and not _is_heading(p):
            add(m.group(1), "inline_label", i); continue
        if _is_heading(p) and (p.style.name or "") != "Title":
            add(t, "heading", i)
    return fields

def load(data: bytes):
    return Document(io.BytesIO(data))

def _set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]: r.text = ""
    else:
        p.add_run(text)

def _insert_after(par, text, style_src=None):
    new = copy.deepcopy(par._p) if style_src is None else copy.deepcopy(style_src._p)
    par._p.addnext(new)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new, par._parent)
    for r in np_.runs[1:]: r._r.getparent().remove(r._r)
    if np_.runs:
        np_.runs[0].text = text; np_.runs[0].bold = None
    else:
        np_.add_run(text)
    return np_

def populate(data: bytes, fields, content: dict) -> bytes:
    doc = load(data)
    paras = list(doc.paragraphs)
    for f in fields:
        txt = (content.get(f["label"]) or "").strip()
        if not txt: continue
        k, ref = f["kind"], f["ref"]
        if k in ("table_cell", "table_below"):
            ti, ri, ci = ref; _set_cell(doc.tables[ti].rows[ri].cells[ci], txt)
        elif k == "cell_para":
            ti, ri, ci, pi = ref; par = doc.tables[ti].rows[ri].cells[ci].paragraphs[pi]
            last = par
            for ln in txt.split("\n")[::-1]:
                pass
            for ln in txt.split("\n"):
                last = _insert_after(last, ln); 
                for r in last.runs: r.bold = False
        elif k == "inline_label":
            p = paras[ref]; r = p.add_run(" " + txt); r.bold = False
        elif k == "heading":
            p = paras[ref]
            # use next empty paragraph as body if present, else insert Normal-styled paragraph
            nxt = paras[ref + 1] if ref + 1 < len(paras) else None
            lines = txt.split("\n")
            if nxt is not None and not nxt.text.strip():
                nxt.text = lines[0]; last = nxt
            else:
                last = p._parent.add_paragraph if False else None
                newp = _insert_after(p, lines[0]); newp.style = doc.styles["Normal"]; last = newp
            for ln in lines[1:]:
                last = _insert_after(last, ln)
    out = io.BytesIO(); doc.save(out); return out.getvalue()

def doc_text(data: bytes) -> str:
    d = load(data)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(parts)
